import hashlib
import io
import re
import zipfile
from pathlib import PurePath
import pandas as pd
import pymupdf
from docx import Document
from .models import Source, Segment

MAX_BYTES = 10 * 1024 * 1024
MAX_CHARS = 120_000

def parse_file(filename: str, data: bytes, **metadata) -> Source:
    if not filename or len(filename) > 180 or re.search(r'[\\/:\x00-\x1f]', filename) or filename in ('.', '..'):
        raise ValueError('文件名无效：不得包含路径或控制字符。')
    ext = PurePath(filename).suffix.lower()
    if ext not in {'.pdf', '.docx', '.txt', '.csv'}:
        raise ValueError('仅支持 PDF、DOCX、TXT、CSV。')
    if not data or len(data) > MAX_BYTES:
        raise ValueError('文件为空或超过 10 MB 限制。')
    segments, warnings = [], []
    try:
        if ext == '.pdf':
            if not data.startswith(b'%PDF-'):
                raise ValueError('PDF 文件签名不符。')
            with pymupdf.open(stream=data, filetype='pdf') as doc:
                if doc.needs_pass or len(doc) > 200:
                    raise ValueError('加密 PDF 或超过 200 页，不支持处理。')
                for i, page in enumerate(doc, 1):
                    text = page.get_text().strip()
                    if text:
                        segments.append(Segment(locator=f'页 {i}', text=text))
                    else:
                        warnings.append(f'页 {i} 无可提取文字，未处理（不支持 OCR）')
        elif ext == '.docx':
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValueError('DOCX 解压体积超过限制。')
                if any('vbaProject' in x for x in archive.namelist()):
                    raise ValueError('不接收含宏文档。')
            doc = Document(io.BytesIO(data))
            segments = [Segment(locator=f'段落 {i}', text=p.text) for i, p in enumerate(doc.paragraphs, 1) if p.text.strip()]
            for ti, table in enumerate(doc.tables, 1):
                for ri, row in enumerate(table.rows, 1):
                    segments.append(Segment(locator=f'表 {ti} 行 {ri}', text=' | '.join(c.text for c in row.cells)))
            warnings.append('不含页眉页脚、文本框及图片内容')
        else:
            text = data.decode('utf-8-sig')
            if '\x00' in text:
                raise ValueError('文本含二进制字符。')
            if ext == '.csv':
                frame = pd.read_csv(io.StringIO(text), dtype=str, keep_default_na=False)
                segments = [Segment(locator=f'数据行 {i + 2}（含表头；多行字段按记录计）', text=' | '.join(f'{k}: {v}' for k, v in row.items())) for i, row in frame.iterrows()]
            else:
                segments = [Segment(locator=f'段落 {i}', text=p) for i, p in enumerate(text.splitlines(), 1) if p.strip()]
    except ValueError:
        raise
    except Exception:
        raise ValueError('无法解析文件，请检查格式；TXT/CSV 需为 UTF-8。') from None
    if not segments:
        raise ValueError('未提取到文字；扫描 PDF 第一版不支持 OCR。')
    if sum(len(s.text) for s in segments) > MAX_CHARS:
        raise ValueError('文本超过 120,000 字符，请拆分后导入；未静默截断。')
    return Source(filename=filename, sha256=hashlib.sha256(data).hexdigest(), title=metadata.pop('title', '') or filename, segments=segments, coverage='全文本；' + '；'.join(warnings) if warnings else '全文本提取（不含图片、嵌入对象）', **metadata)

def citation_valid(fact, sources):
    norm = lambda s: re.sub(r'\s+', '', s)
    return bool(norm(fact.quote)) and any(s.source_id == fact.source_id and any(p.locator == fact.locator and norm(fact.quote) in norm(p.text) for p in s.segments) for s in sources)
